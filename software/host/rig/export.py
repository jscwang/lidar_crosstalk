"""One Excel workbook per session: the raw readings, the run and block tables, the analysis tables and the figures.

`rig export <session folder>` writes `<session folder>/<session name>.xlsx`. The analysis folder is made first if it is
missing. The workbook is built in openpyxl's write-only mode, so a session with a few hundred thousand readings stays
light on memory.

Author: Jessica
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AUTHOR = "Jessica"
MAX_ROWS = 1_048_575            # Excel's sheet limit, less the header row

# (sheet name, file relative to the session folder, what it holds)
TABLES = [
    ("Crossover", "analysis/crossover.csv",
     "Crossover angle per distance and ROI: where the extra ambient falls to 10% of its value at dtheta = 0, "
     "with the 95% bootstrap interval and the predicted angle."),
    ("Effects summary", "analysis/effects_summary.csv",
     "Crosstalk effect at every dtheta, blocks pooled over both sensors: mean, 95% interval, effect size, paired tests."),
    ("Effects by sensor", "analysis/effects_by_sensor.csv", "The same, for sensor A and sensor B separately."),
    ("Pair effects", "analysis/pair_effects.csv", "Per-block paired differences between the AB run and the solo runs."),
    ("Run stats", "analysis/run_stats.csv", "Statistics of every run: ambient, range, dropouts, drift."),
    ("Runs", "runs.csv", "One row per run, with its block, accepted or not, any reject reason, and the number of readings."),
    ("Blocks", "blocks.csv", "One row per block (distance, angle, ROI, repeat): solo ambient levels and the acceptance verdict."),
]
READINGS = ("Readings", "readings.csv",
            "Every reading: sensor, condition, distance, dtheta, ROI, range, signal and ambient rate, SPADs, timestamps.")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")


def _sig(x):
    """Six significant digits: keeps the tables readable without cutting a small p-value to zero."""
    if isinstance(x, float):
        return float(f"{x:.6g}")
    return x


def _clean(v):
    if v is None or (isinstance(v, float) and v != v):
        return None
    if hasattr(v, "item"):          # numpy scalar
        v = v.item()
    if isinstance(v, float):
        return _sig(v)
    return v


def _header(ws, names) -> list:
    cells = []
    for n in names:
        c = WriteOnlyCell(ws, value=str(n))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cells.append(c)
    return cells


def _width(name: str, sample) -> float:
    longest = max([len(str(name))] + [len(str(v)) for v in sample if v is not None])
    return min(max(10, longest + 2), 40)


def _write_frame(wb: Workbook, title: str, df: pd.DataFrame, sig: bool = True) -> None:
    ws = wb.create_sheet(title)
    sample = df.head(200)
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = _width(col, sample[col].tolist())
    ws.freeze_panes = "A2"
    ws.append(_header(ws, df.columns))
    obj = df.astype(object).where(df.notna(), None)
    for row in obj.itertuples(index=False, name=None):
        ws.append([_clean(v) if sig else (None if v is None else (v.item() if hasattr(v, "item") else v)) for v in row])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def _write_lines(wb: Workbook, title: str, lines: list[str], width: int = 120) -> None:
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = width
    for ln in lines:
        ws.append([ln])


def _config_lines(info: dict) -> list[str]:
    cfg = info.get("config") if isinstance(info.get("config"), dict) else info
    if not cfg:
        return []
    def show(key):
        v = cfg.get(key)
        return ", ".join(str(x) for x in v) if isinstance(v, list) else v
    return [f"Config: {cfg.get('name', '')}",
            f"Distances (m): {show('distances_m')}    ROI: {show('rois')}    Repeats: {cfg.get('repeats')}    "
            f"Run length: {(cfg.get('run_ms') or 0) / 1000:g} s",
            f"Angle per side (deg): {show('theta_per_side_deg')}    Timing budget: {cfg.get('tb_ms')} ms    "
            f"Inter-measurement: {cfg.get('imp_ms')} ms    Mode: {cfg.get('mode')}"]


def _neutralise(path: Path) -> None:
    """The writer stamps docProps/app.xml with its own name and version; leave the application field empty."""
    import os
    import zipfile
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"/>')
            zout.writestr(item, data)
    os.replace(tmp, path)


def export_session(session: Path | str, out: Optional[Path | str] = None, n_boot: int = 2000) -> Path:
    """Writes the workbook and returns its path."""
    session = Path(session)
    if not (session / "readings.csv").exists():
        raise FileNotFoundError(f"{session} has no readings.csv: give the session folder (the one rig session made)")
    if not (session / "analysis" / "summary.md").exists():
        from .analysis import analyze
        analyze(session, n_boot=n_boot)

    info = json.loads((session / "session.json").read_text(encoding="utf-8")) if (session / "session.json").exists() else {}
    out = Path(out) if out else session / f"{session.name}.xlsx"

    wb = Workbook(write_only=True)
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.title = f"Crosstalk session {session.name}"

    runs = pd.read_csv(session / "runs.csv")
    blocks = pd.read_csv(session / "blocks.csv")
    n_read = sum(1 for _ in open(session / "readings.csv", "rb")) - 1
    accepted = int(runs["accepted"].astype(str).str.lower().eq("true").sum()) if "accepted" in runs else 0
    acc_blocks = int(blocks["accepted"].astype(str).str.lower().eq("true").sum()) if "accepted" in blocks else 0

    readme = [
        f"Session: {session.name}",
        f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        *_config_lines(info),
        "",
        f"Readings: {n_read}    Runs: {len(runs)} ({accepted} accepted)    Blocks: {len(blocks)} ({acc_blocks} accepted)",
        "",
        "Sheets",
    ]
    for name, rel, what in TABLES + [READINGS]:
        if (session / rel).exists():
            readme.append(f"  {name}: {what}")
    readme += ["  Summary: the written summary of the analysis.",
               "  Figures: the plots of the analysis.",
               "",
               "Notes",
               "  dtheta is the angle between the two sensors' axes (twice the angle turned on each side when both are turned).",
               "  Extra ambient is the ambient rate of the AB run minus the solo ambient rate, in kcps.",
               "  The crossover is the angle where the extra ambient has fallen to 10% of its value at dtheta = 0.",
               "  The Readings sheet is the raw data. A filter on the header row picks one run, sensor or condition."]
    _write_lines(wb, "Read me", readme)

    summary = session / "analysis" / "summary.md"
    if summary.exists():
        _write_lines(wb, "Summary", summary.read_text(encoding="utf-8").splitlines(), width=160)

    for name, rel, _ in TABLES:
        path = session / rel
        if path.exists():
            _write_frame(wb, name, pd.read_csv(path))

    figs = sorted((session / "analysis" / "figures").glob("*.png"))
    if figs:
        ws = wb.create_sheet("Figures")
        row = 1
        for f in figs:
            img = Image(str(f))
            scale = min(1.0, 900.0 / img.width)
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            img.anchor = f"A{row}"
            ws.add_image(img)
            row += int(img.height / 20) + 3

    readings = pd.read_csv(session / READINGS[1])
    if len(readings) > MAX_ROWS:
        readings = readings.head(MAX_ROWS)
        readme_note = f"Only the first {MAX_ROWS} readings fit on one sheet: use readings.csv for the rest."
        _write_lines(wb, "Readings note", [readme_note])
    _write_frame(wb, READINGS[0], readings, sig=False)

    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
        _neutralise(out)
    except PermissionError:
        raise PermissionError(f"cannot write {out}: if the workbook is open in Excel, close it and run the command again")
    return out
