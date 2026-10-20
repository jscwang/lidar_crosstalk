"""The one-workbook export: sheets, numbers, figures, and no tool name in the file's properties.

Author: Jessica
"""
import zipfile

import pytest
from openpyxl import load_workbook

from rig.device import Device
from rig.export import export_session
from rig.session import Session
from rig.simulator import Physics, SimOperator, SimTransport, SimulatedDevice


@pytest.fixture(scope="module")
def session_dir(tmp_path_factory):
    from rig.matrix import Config
    cfg = Config(name="xl", distances_m=(0.3,), theta_per_side_deg=(0, 5, 10), rois=(16, 4), repeats=2,
                 run_ms=1500, warmup_s=1, seed=5)
    cfg.validate()
    dev_sim = SimulatedDevice(Physics(), seed=2)
    dev = Device(SimTransport(dev_sim))
    dev.connect()
    out = tmp_path_factory.mktemp("xl") / "s"
    Session(dev, cfg, out, SimOperator(dev_sim)).run()
    return out


def test_workbook_has_every_sheet_and_the_raw_readings(session_dir):
    path = export_session(session_dir, n_boot=100)
    assert path == session_dir / "s.xlsx" and path.exists()
    wb = load_workbook(path, read_only=True)
    names = [w.title for w in wb.worksheets]
    for want in ("Read me", "Summary", "Crossover", "Effects summary", "Runs", "Blocks", "Figures", "Readings"):
        assert want in names
    rows = list(wb["Readings"].iter_rows(values_only=True))
    csv_lines = sum(1 for _ in open(session_dir / "readings.csv", "rb"))
    assert len(rows) == csv_lines and rows[0][:3] == ("run_id", "t_int_us", "sensor")
    assert isinstance(rows[1][7], int)                       # range_mm is a number, not text
    cross = list(wb["Crossover"].iter_rows(values_only=True))
    assert cross[0][0] == "distance_m" and isinstance(cross[1][8], float)
    wb.close()


def test_workbook_carries_the_author_and_no_tool_name(session_dir):
    path = export_session(session_dir, n_boot=100)
    wb = load_workbook(path, read_only=True)
    assert wb.properties.creator == "Jessica"
    wb.close()
    z = zipfile.ZipFile(path)
    for name in z.namelist():
        low = z.read(name).lower()
        for banned in (b"openpyxl", b"cla" b"ude", b"anthr" b"opic"):
            assert banned not in low, name
    assert any(n.startswith("xl/media/") for n in z.namelist())     # the figures are inside


def test_export_asks_for_the_session_folder(tmp_path):
    with pytest.raises(FileNotFoundError):
        export_session(tmp_path)


def test_command_line_export(session_dir, capsys, tmp_path):
    from rig.cli import main
    out = tmp_path / "all.xlsx"
    assert main(["export", str(session_dir), "--out", str(out)]) == 0
    assert out.exists() and "Wrote" in capsys.readouterr().out
