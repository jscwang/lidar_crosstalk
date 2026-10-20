"""Fill the official IEEE Sensors Letters Word template (sensors_letters_template.docx, downloaded from
https://template-selector.ieee.org/) with the manuscript content, in place of its instructional sample text.

Every paragraph style used below (Subject/Title/Authors/Affiliations/Abstract/Index/Text/Caption/
Heading Refs/References L-SENS) is one already defined in the template; nothing is invented. The template's
own front-matter section (single column: subject, title, authors, affiliations, abstract, index terms) is
edited in place, paragraph by paragraph, so its section break into the two-column body (carried by the
template's paragraph 19) is never touched. Everything from "INTRODUCTION" onward is removed and rebuilt
inside that same two-column section.

Per the template's own "Supplementary Information" paragraph: "Authors may not include any supplementary
information. The journal's rapid-publication model requires short articles ... and no supplementary material
is allowed." The letter's supplement (Algorithm S1, Table S1, Fig. S1) is therefore NOT included in this
file; see MANUSCRIPT_PROGRESS.md for where that content now lives instead.

Run:   python manuscript_ieee/scripts/fill_sensors_letters_docx.py
In:    manuscript_ieee/sensors_letters_template.docx   (as downloaded; not modified)
Out:   manuscript_ieee/main_SensorsLetters.docx

Author: Jessica
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "sensors_letters_template.docx"
FIG = ROOT / "figures"
OUT = ROOT / "main_SensorsLetters.docx"
EQDIR = Path(__file__).resolve().parent / "_eq_png"  # created on first use in render_eq()
COLW = 3.455  # in, single-column width in the template's two-column section (measured from its sectPr)
# Layout switches. The defaults give the journal Word file. make_preprint_docx.py overrides them for the plain,
# single-column mark-up copy of the comprehensive (preprint) version, which also adds the supplement's content.
LAYOUT = {"two_col": True, "preprint": False, "col_w": COLW, "full_w": 7.0, "fig1_w": min(COLW - 0.1, 3.2),
          "table_w": 7.1}

# ----------------------------------------------------------------------- reference numbering, as in main.pdf
CITE_NUM = {
    "ds12385": 1, "stforum": 2, "wermke2019": 3, "whyte2010": 4, "li2015": 5, "kim2015a": 6, "kim2015b": 7,
    "popko2020": 8, "ogurcovs2025": 9, "fasolino2025": 10, "tca9548a": 11, "um2510": 12, "fritsch1980": 13,
    "efron1993": 14,
}
SEC_NUM = {"sec:intro": "", "sec:method": "", "sec:results": "", "sec:discussion": "", "sec:conclusion": ""}
EQ_NUM = {"eq:extra": "1", "eq:pool": "2", "eq:tilted": "3", "eq:geom": "4"}
FIG_NUM = {"fig:setup": "1", "fig:results": "2", "fig:methods": "3"}   # fig:methods: preprint version only
TAB_NUM = {"tab:cross": "I", "tab:divider": "II"}                      # tab:divider: preprint version only
ALG_NUM = {"alg:cross": "1"}                                           # preprint version only

REFERENCES = [
    "STMicroelectronics, VL53L1X: A new generation, long distance ranging time-of-flight sensor based on "
    "ST\u2019s FlightSense technology, Aug. 2024, datasheet DS12385, Rev. 8. Accessed: Oct. 11, 2026. [Online]. "
    "Available: https://www.st.com/resource/en/datasheet/vl53l1x.pdf",
    "J. E. Kvam, \u201cReply to \u2018VCSEL signals interferences for closed sensors (VL53L1X)\u2019,\u201d "
    "STMicroelectronics Community forum, Apr. 2021, accessed: Oct. 11, 2026. [Online]. Available: "
    "https://community.st.com/imaging-sensors-49/vcsel-signals-interferences-for-closed-sensors-vl53l1x-25089",
    "F. Wermke and B. Meffert, \u201cInterference model of two time-of-flight cameras,\u201d in Proc. IEEE "
    "SENSORS, Montreal, QC, Canada, Oct. 2019, pp. 1\u20134.",
    "R. Z. Whyte, A. D. Payne, A. A. Dorrington, and M. J. Cree, \u201cMultiple range imaging camera operation "
    "with minimal performance impact,\u201d in Proc. SPIE, Image Processing: Machine Vision Applications III, "
    "vol. 7538, Feb. 2010, p. 75380I.",
    "L. Li, S. Xiang, Y. Yang, and L. Yu, \u201cMulti-camera interference cancellation of time-of-flight (TOF) "
    "cameras,\u201d in Proc. IEEE Int. Conf. Image Process. (ICIP), Sep. 2015, pp. 556\u2013560.",
    "G. Kim, J. Eom, and Y. Park, \u201cAn experiment of mutual interference between automotive LIDAR "
    "scanners,\u201d in Proc. 12th Int. Conf. Inf. Technol.\u2014New Gener. (ITNG), Apr. 2015, pp. 680\u2013685.",
    "G. Kim, J. Eom, and Y. Park, \u201cInvestigation on the occurrence of mutual interference between pulsed "
    "terrestrial LIDAR scanners,\u201d in Proc. IEEE Intell. Veh. Symp. (IV), Jun. 2015, pp. 437\u2013442.",
    "G. B. Popko, T. K. Gaylord, and C. R. Valenta, \u201cInterference measurements between single-beam, "
    "mechanical scanning, time-of-flight lidars,\u201d Opt. Eng., vol. 59, no. 5, May 2020, art. no. 053106.",
    "A. Ogurcovs, I. Aulika, S. Cartiel, J. Garcia-Pueyo, and A. Mu\u00f1oz, \u201cCrosstalk effects in a dual "
    "ToF-based tactile\u2013proximity sensing platform integrated in a flat PMMA light guide,\u201d Sensors, "
    "vol. 25, no. 23, p. 7319, Dec. 2025.",
    "A. Fasolino, L. E. Andrada, R. Liguori, L. Di Benedetto, A. Rubino, D. Pau, and G. D. Licciardo, "
    "\u201cReal-time crosstalk compensation in multizone time-of-flight sensors,\u201d in Proc. SPIE, "
    "Real-time Processing of Image, Depth, and Video Information 2025, vol. 13526, May 2025, p. 135260D.",
    "Texas Instruments, TCA9548A low-voltage 8-channel I2C switch with reset, Nov. 2019, datasheet SCPS207G. "
    "Accessed: Oct. 3, 2026. [Online]. Available: https://www.ti.com/lit/gpn/TCA9548A",
    "STMicroelectronics, A guide to using the VL53L1X ultra lite driver, Dec. 2025, user manual UM2510, "
    "Rev. 7. Accessed: Oct. 3, 2026. [Online]. Available: "
    "https://www.st.com/resource/en/user_manual/um2510-a-guide-to-using-the-vl53l1x-ultra-lite-driver-"
    "stmicroelectronics.pdf",
    "F. N. Fritsch and R. E. Carlson, \u201cMonotone piecewise cubic interpolation,\u201d SIAM J. Numer. Anal., "
    "vol. 17, no. 2, pp. 238\u2013246, 1980.",
    "B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. New York, NY, USA: Chapman & Hall, 1993.",
]

# ----------------------------------------------------------------------- LaTeX -> plain text (same vocabulary
# the manuscript's .tex sources use; see sections/*.tex)
PREFIX = {"milli": "m", "kilo": "k", "centi": "c", "nano": "n", "micro": "\u00b5"}
UNIT = {"meter": "m", "second": "s", "hertz": "Hz", "degree": "\u00b0", "percent": "%"}


def unit_symbol(u: str) -> str:
    toks = [t.lstrip("\\") for t in re.findall(r"\\[a-zA-Z]+|[a-zA-Z]+", u)]
    if len(toks) == 1 and toks[0] not in UNIT and toks[0] not in PREFIX:
        return toks[0]
    return "".join(PREFIX.get(t, UNIT.get(t, t)) for t in toks)


def _space(u: str) -> str:
    return "" if u in ("\u00b0", "%") else "\u00a0"


def si_sub(text: str) -> str:
    def sirange(m):
        u = unit_symbol(m.group(3))
        return f"{m.group(1)}\u2013{m.group(2)}{_space(u)}{u}"

    def si(m):
        u = unit_symbol(m.group(2))
        return f"{m.group(1)}{_space(u)}{u}"

    text = re.sub(r"\\SIrange\{([^}]*)\}\{([^}]*)\}\{([^}]*)\}", sirange, text)
    text = re.sub(r"\\SI\{([^}]*)\}\{([^}]*)\}", si, text)
    text = re.sub(r"\\numrange\{([^}]*)\}\{([^}]*)\}", lambda m: f"{m.group(1)}\u2013{m.group(2)}", text)
    text = re.sub(r"\\si\{([^}]*)\}", lambda m: unit_symbol(m.group(1)), text)
    text = re.sub(r"\\num\{([^}]*)\}", lambda m: m.group(1), text)
    return text


def cite_sub(text: str) -> str:
    def rep(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = sorted(CITE_NUM[k] for k in keys)
        return ", ".join(f"[{n}]" for n in nums)

    return re.sub(r"~?\\cite\{([^}]*)\}", rep, text)


def xref_sub(text: str) -> str:
    text = re.sub(r"\\eqref\{([^}]*)\}", lambda m: f"({EQ_NUM[m.group(1)]})", text)

    def ref(m):
        key = m.group(1)
        for d in (SEC_NUM, FIG_NUM, TAB_NUM, ALG_NUM):
            if key in d:
                return d[key]
        return "?" + key + "?"

    return re.sub(r"\\ref\{([^}]*)\}", ref, text)


SUPER = {"0": "\u2070", "1": "\u00b9", "2": "\u00b2", "3": "\u00b3", "4": "\u2074", "5": "\u2075",
         "6": "\u2076", "7": "\u2077", "8": "\u2078", "9": "\u2079", "+": "\u207a", "-": "\u207b"}


def to_super(token: str) -> str:
    return "".join(SUPER.get(c, c) for c in token)


def mathtext_sub(s: str) -> str:
    s = s.strip()
    repl = [
        (r"\\dth", "\u0394\u03b8"), (r"\\Delta\\theta", "\u0394\u03b8"), (r"\\theta", "\u03b8"),
        (r"\\varphi", "\u03c6"), (r"\\alpha", "\u03b1"), (r"\\beta", "\u03b2"),
        (r"\\times", "\u00d7"), (r"\\pm", "\u00b1"), (r"\\in\b", "\u2208"), (r"\\sqrt2", "\u221a2"),
        (r"\\mathrm\{([^}]*)\}", r"\1"), (r"\\overline\{\\Delta a\}", "\u0394a\u0305"),
        (r"\\overline\{([^}]*)\}", lambda m: m.group(1) + "\u0305"),
        (r"\\bar\s*\{([A-Za-z])\}", lambda m: m.group(1) + "\u0304"),
        (r"\\bar\s+([A-Za-z])", lambda m: m.group(1) + "\u0304"),
        (r"\\tan\b", "tan"), (r"\\exp\b", "exp"), (r"\\dots", "\u2026"),
        (r"\^\\circ", "\u00b0"), (r"\\circ", "\u00b0"),
    ]
    for pat, rep in repl:
        s = re.sub(pat, rep, s)
    s = re.sub(r"_\{([^}]*)\}", lambda m: "\u2009" + m.group(1), s)
    s = re.sub(r"_([A-Za-z0-9])", lambda m: "\u2009" + m.group(1), s)
    s = re.sub(r"\^\{([^}]*)\}", lambda m: to_super(m.group(1)), s)
    s = re.sub(r"\^([A-Za-z0-9])", lambda m: to_super(m.group(1)), s)
    s = s.replace("\\{", "{").replace("\\}", "}")
    s = re.sub(r"\\[A-Za-z]+\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\([A-Za-z]+)", r"\1", s)
    return s


def plain_sub(text: str) -> str:
    text = si_sub(text)
    text = cite_sub(text)
    text = xref_sub(text)
    text = re.sub(r"\$([^$]+)\$", lambda m: mathtext_sub(m.group(1)), text)
    text = text.replace(r"\,", "\u2009").replace("~", "\u00a0")
    text = text.replace("---", "\u2014").replace("--", "\u2013")
    text = text.replace("``", "\u201c").replace("''", "\u201d")
    text = text.replace(r"\%", "%")
    text = re.sub(r"\\label\{[^}]*\}", "", text)
    return text


EMPH_RE = re.compile(r"\\emph\{([^}]*)\}")


def add_rich(paragraph, raw_text: str):
    text = plain_sub(raw_text)
    pos = 0
    for m in EMPH_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        r = paragraph.add_run(m.group(1))
        r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


# ----------------------------------------------------------------------- equation images
def render_eq(mathtext: str, name: str, fontsize=14):
    EQDIR.mkdir(exist_ok=True)
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["mathtext.fontset"] = "cm"
    fig = plt.figure(figsize=(0.01, 0.01))
    t = fig.text(0, 0, f"${mathtext}$", fontsize=fontsize)
    fig.canvas.draw()
    bbox = t.get_window_extent()
    w, h = bbox.width / fig.dpi + 0.1, bbox.height / fig.dpi + 0.08
    plt.close(fig)
    fig = plt.figure(figsize=(w, h))
    fig.text(0.02, 0.5, f"${mathtext}$", fontsize=fontsize, va="center")
    out = EQDIR / f"{name}.png"
    fig.savefig(out, dpi=220, transparent=True)
    plt.close(fig)
    return out, w, h


EQUATIONS = {
    "1": r"\Delta a_s(b)=\bar{a}_s^{AB}(b)-\bar{a}_s^{solo}(b)",
    "2": r"\Delta a(b)=\frac{1}{2}\left[\Delta a_A(b)+\Delta a_B(b)\right]",
    "3": r"g(\Delta\theta)=\frac{a\,(1+b\,\Delta\theta)}{1+\exp[(\Delta\theta-x_0)/w]}",
    "4": r"s+2d\,\tan(\Delta\theta_g/2)=d\,[\tan(E/2)+\tan(F/2)]",
}


def add_equation(doc, styles, num: str):
    img, w, h = render_eq(EQUATIONS[num], f"eq{num}", fontsize=13)
    COLW = LAYOUT["col_w"]
    max_w = COLW - 0.5
    if w > max_w:
        h *= max_w / w
        w = max_w
    p = doc.add_paragraph(style=styles["Text L-SENS"])
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLW / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLW), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    run = p.add_run()
    run.add_picture(str(img), width=Inches(w))
    p.add_run(f"\t({num})")
    return p


def section_break(doc, cols):
    """End the section written so far with a continuous section break carrying `cols` columns.

    The new paragraph holds a copy of the body's own <w:sectPr> (the template's two-column settings);
    only the column count is changed. Writing continues in the body-level sectPr, i.e. the two-column
    section, so call this once to leave the two-column stretch (cols=2 describes the stretch that ends
    here) and once more to leave the one-column stretch (cols=1)."""
    import copy
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    sp = copy.deepcopy(doc.element.body.find(qn("w:sectPr")))
    cols_el = sp.find(qn("w:cols"))
    for k in list(cols_el.attrib):
        del cols_el.attrib[k]
    if cols == 2:
        cols_el.set(qn("w:num"), "2")
        cols_el.set(qn("w:space"), "360")
    else:
        cols_el.set(qn("w:space"), "720")
    t = sp.find(qn("w:titlePg"))
    if t is not None:
        sp.remove(t)
    p._p.get_or_add_pPr().append(sp)
    return p


def shade_cell(cell, hex_color="FFFFFF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


# ----------------------------------------------------------------------- front matter (single-column section)
def fill_front_matter(doc):
    p = doc.paragraphs
    # 0: Subject L-SENS
    set_text(p[0], "Crosstalk and Interference in Time-of-Flight Sensor Arrays")
    # 1,3: Title L-SENS spacer (leave empty); 2: Title L-SENS title
    set_text(p[2], "Angle-Dependent Crosstalk Between Two Adjacent VL53L1X Time-of-Flight Sensors")
    # 4: Authors L-SENS
    set_text(p[4], "[Author names, order and affiliations to be supplied before submission]")
    # 5-10: Affiliations L-SENS (sample has 6 lines: 4 affiliations + 2 member-grade notes) -> collapse to one
    # placeholder and delete the rest, since none of them carries the section break (only paragraph 19 does)
    set_text(p[5], "[Affiliation(s) to be supplied before submission]")
    for extra in p[6:11]:
        delete_paragraph(extra)
    # re-fetch after deletion
    p = doc.paragraphs
    idx = {id(par._p): i for i, par in enumerate(p)}
    # locate by style name from here on, more robust than fixed indices after the deletions above
    by_style = {}
    for par in p:
        by_style.setdefault(par.style.name, []).append(par)
    # the "Received/revised ... Dates will be inserted by IEEE" line: author-facing text does not belong here
    for par in by_style.get("Index L-SENS", []):
        if "Dates will be inserted" in par.text:
            set_text(par, "")
    # Abstract L-SENS
    abstract = (
        "Abstract\u2014Time-of-flight ranging sensors mounted side by side can disturb one another, but we "
        "found no published measurement of how this depends on the angle between them. We measured the extra "
        "background light that appears in one VL53L1X sensor when a second, identical sensor 20 millimeters "
        "away is also ranging, as the two were turned apart by up to 24 degrees. One pair of sensors was "
        "tested against a grey board at 0.3, 0.6 and 0.9 meters in two window sizes, with five repeats per "
        "setting and a divider test. The extra light fell to 10 percent of its peak between 16 and 18 degrees "
        "for the full array and between 8 and 12 degrees for the window, and to 1 percent at about 20 to 22 "
        "and 11 to 13 degrees; these angles grew with distance. Methods for locating the 10 percent angle "
        "disagreed by up to one degree, ten to forty-four times the sampling uncertainty, so we report a "
        "bracket and an interpolated value. A geometric model using the datasheet field of view overestimated "
        "the angle at which the effect disappears by three to seven degrees, and no single emitter beam angle "
        "fitted both windows. A thin divider changed the effect by no more than about four percent in ten of "
        "twelve settings, which does not show that direct coupling is absent. Range readings were barely "
        "affected. Turning each sensor outward by about 11 degrees reduced the effect below 1 percent of its "
        "peak for the full array."
    )
    for par in by_style.get("Abstract L-SENS", []):
        set_text(par, abstract)
    # Index Terms (second 'Index L-SENS' paragraph)
    for par in by_style.get("Index L-SENS", []):
        if par.text == "" and par is not None:
            pass
    index_terms_par = [par for par in doc.paragraphs if par.style.name == "Index L-SENS"]
    # after clearing the dates line, the remaining non-empty-by-design Index L-SENS para is Index Terms;
    # find it by matching leftover sample text instead of relying on clearing order
    for par in index_terms_par:
        if "Index Terms" in par.text or "key words" in par.text:
            set_text(par, "Index Terms\u2014Crosstalk, interference, sensor arrays, time-of-flight sensors, "
                           "VL53L1X.")


def set_text(paragraph, text):
    """Replace a paragraph's content outright: the template's sample paragraphs carry leftover
    <w:hyperlink> runs and <w:proofErr> spell-check markers that paragraph.runs does not see (it only
    returns direct <w:r> children), so clearing .runs alone leaves hyperlink text glued onto the end.
    Strip every child except <w:pPr> (which carries the style/formatting) and add one clean run."""
    p = paragraph._p
    pPr = p.find(qn("w:pPr"))
    for child in list(p):
        if child is not pPr:
            p.remove(child)
    if text:
        paragraph.add_run(text)


def delete_paragraph(paragraph):
    el = paragraph._p
    el.getparent().remove(el)


# ----------------------------------------------------------------------- body (two-column section)
def clear_body(doc):
    """Remove every paragraph from 'INTRODUCTION' to the end, leaving the section's own sectPr (owned by the
    body element itself, not by any paragraph) untouched."""
    body = doc.element.body
    started = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue  # the final section properties; keep
        if not started:
            texts = child.findall(".//" + qn("w:t"))
            if any((t.text or "").strip() == "INTRODUCTION" for t in texts):
                started = True
        if started:
            body.remove(child)


def build_body(doc, styles):
    H1 = styles["Heading 1"]
    H2 = styles["Heading 2"]
    TXT = styles["Text L-SENS"]
    CAP = styles["Caption L-SENS"]
    HREF = styles["Heading Refs L-SENS"]
    ACK = styles.get("Acknowledgment L-SENS", TXT)
    REF = styles["References L-SENS"]

    def h1(text):
        return doc.add_paragraph(text, style=H1)

    def h2(text):
        return doc.add_paragraph(text, style=H2)

    def body_p(latex_text):
        p = doc.add_paragraph(style=TXT)
        add_rich(p, latex_text)
        return p

    def caption(latex_text):
        p = doc.add_paragraph(style=CAP)
        add_rich(p, latex_text)
        return p

    def figure(path, width_in):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width_in))
        return p

    # -------------------------------------------------------------- INTRODUCTION
    h1("INTRODUCTION")
    body_p(
        "Compact direct time-of-flight (ToF) sensors such as the VL53L1X, which pairs a \\SI{940}{\\nano\\meter} "
        "laser emitter with an array of single-photon avalanche diodes (SPADs) and a programmable region of "
        "interest (ROI)~\\cite{ds12385}, are often mounted in groups to widen coverage. One developer building a "
        "drone with several of them asked on a manufacturer forum whether the units disturb one "
        "another~\\cite{stforum}. Mutual interference has been studied for ToF cameras, where a model of the "
        "distance error was proposed for two cameras~\\cite{wermke2019} and methods for operating several were "
        "developed~\\cite{whyte2010,li2015}, and for scanning lidars, where it was measured~\\cite{kim2015a,kim2015b} "
        "and direct and scattered interference were distinguished~\\cite{popko2020}. Crosstalk between two "
        "multizone ToF sensors viewing through one light guide has been characterised for a tactile and proximity "
        "skin~\\cite{ogurcovs2025}, and crosstalk compensation in multizone ToF sensors has been "
        "reported~\\cite{fasolino2025}. For the VL53L1X, a manufacturer employee replied on that forum that units "
        "generally do not interfere, while relaying a report of interference at \\SI{5}{\\centi\\meter} when two "
        "face each other~\\cite{stforum}. In the sources we searched (web searches of the research literature, "
        "manufacturer documents and community posts, not a bibliographic-database search) we found no published "
        "measurement of how the interaction between two adjacent VL53L1X-class sensors depends on the angle "
        "between them.")
    body_p(
        ("This paper" if LAYOUT["preprint"] else "This letter") + " measures that dependence for one pair of VL53L1X sensors \\SI{20}{\\milli\\meter} apart, in "
        "two ROI sizes and at three distances. Its contributions are: (i) the angle at which the extra ambient "
        "rate falls to \\SI{10}{\\percent} and \\SI{1}{\\percent} of its peak, with a model-free bracket and the "
        "spread between estimators, which is 10 to 44 times the sampling uncertainty; (ii) a non-monotonic "
        "response of the small ROI at two of three distances; (iii) a check of a geometric overlap model against "
        "the datasheet's field-of-view convention; and (iv) a divider test, with its limits stated.")

    # -------------------------------------------------------------- METHODOLOGY
    h1("METHODOLOGY")
    h2("Apparatus and Geometry")
    body_p(
        "Two VL53L1X sensors~\\cite{ds12385} on Pololu 3415 carrier boards (unit~3 at position~A, unit~1 at B) "
        "were read by an ESP32-S3 through a TCA9548A I$^2$C multiplexer~\\cite{tca9548a} at "
        "\\SI{400}{\\kilo\\hertz}. The STMicroelectronics ultra-lite driver~\\cite{um2510} set short distance "
        "mode, a \\SI{33}{\\milli\\second} timing budget and a \\SI{40}{\\milli\\second} period. Each "
        "interrupt-timestamped reading gave the range, its status, and the total signal and ambient rates "
        "(kilocounts per second, kcps, summed over the enabled SPADs, in \\SI{8}{kcps} steps). Two ROIs were "
        "used, the full $16\\times16$ array and a $4\\times4$ window, each centred on the factory optical "
        "centre. The datasheet gives their fields of view as \\SI{27}{\\degree} and \\SI{15}{\\degree}, as "
        "diagonal values, and gives no emitter beam angle~\\cite{ds12385}.")
    body_p(
        "The sensors, nominally $s=\\SI{20}{\\milli\\meter}$ apart, sat on pucks turned by $-\\theta$ (A) and "
        "$+\\theta$ (B), so their axes toed out by $\\dth=2\\theta$ (Fig.~\\ref{fig:setup}). Pins in a plate "
        "with holes at \\SI{1}{\\degree} steps set the angles (accuracy not verified). The sweep covered "
        "$\\theta=\\SI{0}{\\degree}$ and \\SIrange{2}{12}{\\degree}, so $\\dth=0,4,6,\\dots,\\SI{24}{\\degree}$ "
        "(12 levels). The target was a flat dark-grey board \\SI{1.2}{\\meter} wide at tape distances $d=0.3$, "
        "$0.6$ and \\SI{0.9}{\\meter} (the sensors read \\SIrange{9}{30}{\\milli\\meter} less). Before each "
        "setting a guided check required, three times in a row, the mean $16\\times16$ range within "
        "\\SI{15}{\\milli\\meter} of a stored value and B$-$A within \\SI{10}{\\milli\\meter} of a per-distance "
        "offset (2.5, 2.5 and \\SI{3.0}{\\milli\\meter}, from a squared board), and a self-test to pass. Room "
        "temperature was not recorded.")
    figure(FIG / "fig_setup.png", LAYOUT["fig1_w"])
    caption(
        "Fig. 1. Measurement geometry (schematic, not to scale): sensors A and B toe out by $\\dth=2\\theta$; "
        "the common footprint is where B's emitter cone and A's receive field overlap. The divider was present "
        "only in the divider tests. Cone angles are generic.")

    h2("Acquisition")
    body_p(
        "Each distance was one campaign on its own day, angles ascending (resumed once after an interruption at "
        "0.3 and \\SI{0.6}{\\meter}), after a \\SI{600}{\\second} two-sensor warm-up. Each cell (zone $\\times$ "
        "$\\dth$) had five blocks of three \\SI{30}{\\second} runs, A alone, B alone and A with B, in seeded "
        "random order with the idle sensor stopped (\\numrange{748}{766} readings per active sensor). A run was "
        "accepted if it had no abort or I$^2$C/read error, the idle sensor stayed silent, the reading count was "
        "within \\SI{5}{\\percent} of expected and at most \\SI{1}{\\percent} of results were missed. A block was "
        "accepted if its runs were and each sensor's solo mean ambient rate was within \\SI{10}{\\percent} (at "
        "least \\SI{0.5}{kcps}) of that of the previous accepted block at the same distance, angle and zone; "
        "failures were repeated up to three times. Blocks failing three times (1, 6 and 7 of 120 at 0.3, 0.6 and "
        "\\SI{0.9}{\\meter}) were re-measured in separate refill sessions and replaced; excluding them moves no "
        "crossover by more than \\SI{0.01}{\\degree}.")

    h2("Crosstalk Quantification")
    body_p("For sensor $s\\in\\{\\mathrm{A},\\mathrm{B}\\}$ in block $b$, the extra ambient rate is the "
           "difference of run-mean ambient rates $\\bar a$ with both sensors ranging and with $s$ alone, "
           "pooled over the sensors:")
    add_equation(doc, styles, "1")
    add_equation(doc, styles, "2")
    body_p(
        "The block is the unit of analysis; five blocks give a cell mean $\\overline{\\Delta a}$ and standard "
        "deviation. A difference is used, not a ratio to the solo rate, which for the $4\\times4$ window is only "
        "\\SIrange{0.1}{9.4}{kcps} and so dominated by noise. Range effects are the difference of mean ranges "
        "(bias) and the ratio of range standard deviations between two-sensor and solo runs, from valid "
        "readings, pooled as in \\eqref{eq:pool} (geometric mean for the ratio).")

    h2("Crossover and Uncertainty")
    body_p(
        "Let $\\Delta\\theta_{\\mathrm{pk}}$ be the level with the largest cell mean "
        "$\\overline{\\Delta a}_{\\mathrm{pk}}$. The crossover at threshold $\\varphi$ is the first angle beyond "
        "$\\Delta\\theta_{\\mathrm{pk}}$ at which the response falls to $\\varphi\\,\\overline{\\Delta "
        "a}_{\\mathrm{pk}}$ ($\\varphi=0.10$ unless stated). With levels \\SI{2}{\\degree} apart, the data alone "
        "fix only a \\emph{bracket} (the last level above and the first at or below the threshold); a point "
        "estimate needs an interpolant or a fit. The primary estimator is shape-preserving piecewise cubic "
        "Hermite interpolation (PCHIP)~\\cite{fritsch1980} of the cell means. Four checks follow: linear "
        "interpolation; a logistic fitted to all levels, with the threshold taken from its fitted value at "
        "$\\dth=0$ (the project's earlier estimator); the same logistic fitted from the peak outward; and a "
        "tilted logistic that can rise before it falls,")
    add_equation(doc, styles, "3")
    body_p(
        "which is the logistic for $b=0$. Sampling uncertainty is a block bootstrap~\\cite{efron1993}: the five "
        "blocks of every level are resampled with replacement $B=2000$ times, with common resamples for all "
        "estimators, and percentile 95\\,\\% intervals are reported. These describe within-session "
        "repeatability only: each distance was measured once with one sensor pair, so between-session and "
        "between-unit variance are not estimated. The analytical uncertainty is the range across estimators; we "
        "also vary $\\varphi\\in\\{0.05,0.10,0.20\\}$ and analyse A and B separately.")
    if LAYOUT["preprint"]:
        body_p("The five estimators are labelled as follows in the figures and in Algorithm~\\ref{alg:cross}:")
        for item in [
            "M1: three-parameter logistic fitted to all levels; reference is the fitted value at $\\dth=0$ (the "
            "project's original estimator);",
            "M2: the same logistic fitted only from the observed peak level outward; reference is the fitted "
            "value there (identical to M1 when the peak is at $\\dth=0$);",
            "M3: piecewise-linear interpolation of the block means; reference is the observed peak mean; first "
            "crossing after the peak;",
            "M4: PCHIP interpolation of the block means, same reference and crossing rule (the primary "
            "estimator);",
            "M5: the tilted logistic of \\eqref{eq:tilted}; reference is its own maximum; first crossing after "
            "that maximum.",
        ]:
            p = doc.add_paragraph(style=TXT)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.add_run("•  ")
            add_rich(p, item)
        body_p("Algorithm~\\ref{alg:cross} gives the procedure for one cell and one threshold; the implemented "
               "form is scripts/crossover_analysis.py. In steps 3–4 the references of M1, M2 and M5 follow "
               "their definitions above rather than the observed peak $\\bar y_p$.")
        caption("Algorithm 1. Crossover for one cell (zone, distance)")
        add_algorithm_box(doc)

    h2("Geometric Reference")
    body_p("The project's model puts the end of the overlap between B's lit spot and A's receive field at "
           "$\\Delta\\theta_g$, where")
    add_equation(doc, styles, "4")
    body_p(
        "with $E$ the emitter cone and $F$ the receive field. As coded, $E=\\SI{27}{\\degree}$ is assumed and "
        "$F$ is the datasheet's diagonal value. We also evaluate \\eqref{eq:geom} with $F$ converted to an "
        "in-plane value for a square array, $\\tan\\beta=\\tan\\alpha/\\sqrt2$, and invert it for the $E$ that "
        "reproduces the angle at which the response falls to \\SI{1}{\\percent} of its peak. This is a "
        "sensitivity analysis: the orientation of the array axes relative to the toe-out plane was not "
        "recorded.")

    h2("Divider Check")
    body_p(
        "At each distance a thin black strip stood on edge between the sensor windows; the protocol required it "
        "to stay out of both fields of view and within about \\SI{15}{\\milli\\meter} of the sensor chips, but "
        "its size and position were not recorded. The cells at $\\theta=\\SI{2}{\\degree}$ and \\SI{5}{\\degree} "
        "($\\dth=4$ and \\SI{10}{\\degree}) were re-measured in both zones, five blocks each, and compared with "
        "the no-divider cells through the difference of cell means of \\eqref{eq:pool}. At \\SI{0.6}{\\meter} "
        "the divider cells were measured twice, the second time with the strip protruding less; the second set "
        "is used and the first reported as a replicate.")

    # -------------------------------------------------------------- RESULTS
    h1("RESULTS")
    # Fig. 2 and Table I span both columns in the letter (7.2 in wide); in the template's two-column section
    # that needs a one-column stretch between two continuous section breaks.
    if LAYOUT["two_col"]:
        section_break(doc, cols=2)
    figure(FIG / "fig_results.png", LAYOUT["full_w"])
    caption(
        "Fig. 2. (a), (b) Extra ambient rate against the angle between the sensors, as a percentage of each "
        "cell's observed peak, for the $16\\times16$ and $4\\times4$ zones at three distances. Markers are "
        "cell means of five blocks (error bars, one standard deviation, are mostly hidden); lines are PCHIP "
        "interpolants; filled triangles mark the \\SI{10}{\\percent} crossing and open triangles the geometric "
        "end-of-overlap angle as coded. (c) The \\SI{10}{\\percent} crossing (bar: model-free bracket) and the "
        "angle at \\SI{1}{\\percent} of the peak against distance, with the model as coded (solid) and with the "
        "receive field converted to an in-plane value (dashed).")
    caption("Table I. Angle between the sensors (degrees) at which the extra ambient rate falls to "
            "\\SI{10}{\\percent} (crossover) and \\SI{1}{\\percent} of its peak, and the model's end-of-overlap "
            "angle")
    add_table_crossover(doc)
    if LAYOUT["preprint"]:
        figure(FIG / "fig_crossover_methods.png", LAYOUT["full_w"])
        caption(
            "Fig. 3. Extra ambient rate against $\\dth$ in each cell, as a percentage of the largest cell mean. "
            "Points are block means $\\pm$ SD ($n=5$). Curves: M1 logistic, M5 tilted logistic and M4 PCHIP "
            "interpolant. Triangles mark where each of these three crosses \\SI{10}{\\percent} of the peak "
            "(colour identifies the estimator). The dash-dotted vertical line is the geometric end-of-overlap "
            "angle of the model of \\eqref{eq:geom}, with the datasheet field of view as coded.")
    if LAYOUT["two_col"]:
        section_break(doc, cols=1)

    h2("Angular Response")
    body_p(
        "The extra ambient rate peaked at 1006, 665 and \\SI{302}{kcps} for the $16\\times16$ zone and at 273, "
        "65 and \\SI{24}{kcps} for the $4\\times4$ window at 0.3, 0.6 and \\SI{0.9}{\\meter}. For the full array "
        "that is 23, 8 and 3.5 times the solo rate; no ratio is given for the window, whose solo rate is near "
        "the noise floor. The response fell with toe-out at every distance for the full array and at "
        "\\SI{0.3}{\\meter} for the window (Fig.~\\ref{fig:results}a,b). At 0.6 and \\SI{0.9}{\\meter} the "
        "window instead peaked at $\\dth=\\SI{4}{\\degree}$, above its value at \\SI{0}{\\degree} by "
        "\\SI{4.5}{\\percent} (Welch $p=0.0024$) and \\SI{23}{\\percent} ($p<0.0001$); at \\SI{0.3}{\\meter} it "
        "was \\SI{25}{\\percent} lower. Distance and session coincide, so this cannot be assigned to distance; "
        "the maximum lies between \\SI{0}{\\degree} and \\SI{6}{\\degree} ($\\theta=\\SI{1}{\\degree}$ was not "
        "measured).")

    h2("Crossover and Its Uncertainty")
    body_p(
        "The model-free bracket of the \\SI{10}{\\percent} crossing (Table~\\ref{tab:cross}) was "
        "\\SIrange{16}{18}{\\degree} for the full array at all three distances and \\numrange{8}{10}, "
        "\\numrange{8}{10} and \\SIrange{10}{12}{\\degree} for the window. PCHIP put the crossing at 16.15, "
        "16.97 and \\SI{17.90}{\\degree} (full array) and 8.22, 9.47 and \\SI{10.80}{\\degree} (window), rising "
        "with distance in all five estimators." + (
            " Fig.~\\ref{fig:methods} compares the block means, three of the estimators and their "
            "\\SI{10}{\\percent} crossings in every cell; in the $4\\times4$ zone at 0.6 and \\SI{0.9}{\\meter}, "
            "where the response peaks at $\\dth=\\SI{4}{\\degree}$ rather than at \\SI{0}{\\degree}, the three "
            "curves differ most at small angles, near the peak, while their crossings stay close."
            if LAYOUT["preprint"] else "") + " The bootstrap half-width of the PCHIP value is "
        "\\SIrange{0.01}{0.07}{\\degree}, but the five estimators span \\SIrange{0.13}{1.02}{\\degree} within a "
        "cell, 10 to 44 times as much. The logistic of the project's earlier analysis lay "
        "\\SIrange{0.9}{1.0}{\\degree} above PCHIP for the full array, outside the bracket at "
        "\\SI{0.9}{\\meter} (\\SI{18.84}{\\degree} against \\SIrange{16}{18}{\\degree}), and it fitted worse "
        "than a tilted logistic (Akaike difference above 140 in five of six cells). PCHIP and the tilted "
        "logistic agree within \\SI{0.14}{\\degree} in every cell. Doubling or halving the threshold moves the "
        "crossing by \\SIrange{0.7}{1.7}{\\degree}, so the \\SI{10}{\\percent} value is a convention. Analysed "
        "separately, sensor~A crossed over \\SIrange{1.0}{1.1}{\\degree} earlier than B for the full array and "
        "\\SIrange{0.6}{1.0}{\\degree} later for the window, with the same sign at every distance; unit and "
        "path direction are confounded.")

    h2("Angle at 1\u2009% of the Peak")
    body_p(
        "The extra ambient fell to \\SI{1}{\\percent} of its peak at 19.9, 20.9 and \\SI{21.6}{\\degree} (full "
        "array) and 11.4, 12.5 and \\SI{13.1}{\\degree} (window). At those angles the full array's extra "
        "ambient was still 0.19, 0.09 and 0.04 times the solo rate. The range readings were barely affected: "
        "at $\\dth=\\SI{0}{\\degree}$ the range bias was at most \\SI{0.3}{\\milli\\meter} in magnitude, and "
        "the range standard deviation was \\SIrange{7}{10}{\\percent} larger with both sensors ranging for the "
        "full array and \\SI{13}{\\percent} larger for the window at \\SI{0.3}{\\meter}; at \\SI{24}{\\degree} "
        "both were indistinguishable from no effect. The window returned valid ranges for about half of its "
        "readings at \\SI{0.6}{\\meter} and for none at \\SI{0.9}{\\meter}, so range effects are not available "
        "for it there.")

    h2("Geometric Model")
    body_p(
        "As coded, the overlap model of \\eqref{eq:geom} puts the end of overlap at \\SIrange{23.4}{25.8}{"
        "\\degree} (full array) and \\SIrange{17.3}{19.8}{\\degree} (window), which is \\SIrange{3.4}{4.3}{"
        "\\degree} and \\SIrange{5.9}{6.8}{\\degree} beyond the \\SI{1}{\\percent} angles. Converting the "
        "datasheet's diagonal field of view to an in-plane value moves the full-array values to "
        "\\SIrange{19.5}{21.9}{\\degree}, within \\SI{0.5}{\\degree} of the \\SI{1}{\\percent} angles, but "
        "leaves the window \\SIrange{3.8}{4.6}{\\degree} too late (Fig.~\\ref{fig:results}c). Inverting the "
        "model for the emitter cone gives \\SIrange{26}{28}{\\degree} for the full array and "
        "\\SIrange{18}{20}{\\degree} for the window: nearly constant with distance within each zone (spread at "
        "most \\SI{1.8}{\\degree}) but not between zones. The model reproduces how the angle grows with "
        "distance but not its dependence on zone, under either convention.")

    h2("Divider")
    body_p(
        "With the strip between the sensors the extra ambient changed by $-7.7$ to $+17.2$\\,\\% across the "
        "twelve cells (both zones, $\\dth=4$ and \\SI{10}{\\degree}, three distances), within $\\pm4.2$\\,\\% "
        "in ten of them; the exceptions are the window cells at \\SI{10}{\\degree}, where the absolute changes "
        "were $-0.5$ and $+0.8$\\,kcps. The sign followed the distance, not the zone or angle: all four cells "
        "fell at \\SI{0.3}{\\meter} ($-0.7$ to $-7.7$\\,\\%), all rose at \\SI{0.6}{\\meter} ($+3.9$ to "
        "$+17.2$\\,\\%) and they changed by $-0.1$ to $+2.5$\\,\\% at \\SI{0.9}{\\meter}. The "
        "\\SI{95}{\\percent} bootstrap intervals excluded zero in ten cells (all but the \\SI{0.9}{\\meter} "
        "window), yet the sensors' solo ambient rate differed by $-20$ to $+18$\\,\\% between the divider and "
        "no-divider sessions, so session differences are as large as the divider effects. The two "
        "\\SI{0.6}{\\meter} strip placements agreed within \\SI{1.6}{\\percent}."
        + (" Table~\\ref{tab:divider} gives each cell with its bootstrap interval." if LAYOUT["preprint"] else ""))
    if LAYOUT["preprint"]:
        caption("Table II. Change in the extra ambient rate with the divider, \\si{\\percent} of the no-divider "
                "cell mean (\\SI{95}{\\percent} bootstrap interval)")
        add_table_divider(doc)

    # -------------------------------------------------------------- DISCUSSION
    h1("DISCUSSION")
    disc = [
        ("Practical reading.", "For this pair, turning each sensor outward by about \\SI{11}{\\degree} (full "
         "array) or \\SI{6.5}{\\degree} (window) brought the extra ambient rate below \\SI{1}{\\percent} of "
         "its peak at all three distances, and by \\SI{9}{\\degree} and \\SI{5.4}{\\degree} below "
         "\\SI{10}{\\percent} (values at \\SI{0.9}{\\meter}). Because the sensors' own ambient rate differs "
         "between zones and distances, \\SI{1}{\\percent} of the peak is not a uniform criterion: at "
         "\\SI{0.3}{\\meter} it was still about a fifth of the full array's solo rate. The effect shows mainly "
         "in the ambient channel; range bias stayed within \\SI{0.3}{\\milli\\meter}."),
        ("The geometric model.", "The datasheet gives the field of view as a diagonal value and no emitter "
         "beam angle~\\cite{ds12385}, so the model's two inputs are conventions. An in-plane receive field "
         "brought the full array within \\SI{0.5}{\\degree} but not the window, and no single emitter cone "
         "fits both. Candidate explanations were not tested: soft beam edges, a window field that is not a "
         "narrower copy of the full cone, array axes outside the toe-out plane. The window's non-monotonic "
         "response, present at 0.6 and \\SI{0.9}{\\meter} but not at \\SI{0.3}{\\meter}, may reflect near-field "
         "saturation at \\SI{0.3}{\\meter} (\\SI{273}{kcps}) or session differences."),
        ("The divider.", "The strip changed the extra ambient by $-7.7$ to $+17.2$\\,\\% (within $\\pm4.2$\\,\\% "
         "in ten of twelve cells), with a sign set by distance, which points to session differences rather "
         "than a strip effect. That does not show that direct coupling is absent: the strip's size and "
         "position were not recorded, it may block only part of a direct path, and the divider cells came "
         "from other sessions."),
        ("Limits.", "One sensor pair; one session per distance; angle and time confounded; unit and path "
         "direction confounded; $\\theta=\\SI{1}{\\degree}$ unsampled; plate accuracy, board reflectance, room "
         "light and strip geometry unrecorded. Re-measured blocks are at most \\SI{6}{\\percent} of those "
         "accepted, mostly in the tails."),
    ]
    for lead, rest in disc:
        p = doc.add_paragraph(style=TXT)
        r = p.add_run(lead + " ")
        r.italic = True
        add_rich(p, rest)

    # -------------------------------------------------------------- CONCLUSION
    h1("CONCLUSION")
    body_p(
        "For one pair of VL53L1X sensors \\SI{20}{\\milli\\meter} apart, the extra ambient rate from running "
        "both fell to \\SI{10}{\\percent} of its peak at \\SIrange{16}{18}{\\degree} of total toe-out for the "
        "full array and \\SIrange{8}{12}{\\degree} for the $4\\times4$ window, and to \\SI{1}{\\percent} at "
        "about \\SIrange{20}{22}{\\degree} and \\SIrange{11}{13}{\\degree}, growing with distance from 0.3 to "
        "\\SI{0.9}{\\meter}; range readings were barely affected. The choice of estimator, not sampling, "
        "limits the precision of these angles, and a geometric model with the datasheet's field of view "
        "overshoots them by \\SIrange{3}{7}{\\degree}. A thin divider changed the effect by no more than about "
        "\\SI{4}{\\percent} in ten of twelve settings, which does not rule out direct coupling. Other units "
        "and targets, a unit swap and $\\theta=\\SI{1}{\\degree}$ are needed before these angles serve as a "
        "design rule.")

    # -------------------------------------------------------------- DATA AND CODE AVAILABILITY (added; not
    # part of the 2016 template, which predates IEEE's data-availability-statement policy)
    h1("DATA AND CODE AVAILABILITY")
    body_p("[The logged sessions, merged data and all analysis and figure scripts will be deposited in an "
           "archive with a DOI, to be inserted" + ("" if LAYOUT["preprint"] else " before submission") + ".]")

    # -------------------------------------------------------------- ACKNOWLEDGMENT
    doc.add_paragraph("ACKNOWLEDGMENT", style=HREF)
    p = doc.add_paragraph(style=ACK)
    p.add_run("[Funding sources, if any, to be supplied by the authors.] [AI-use disclosure, required by "
              "IEEE and to be completed by the authors: an AI system, ")
    r = p.add_run("name and version to be supplied")
    r.italic = True
    p.add_run(", was used to write and run the analysis scripts, to generate the figures and tables, and to "
              "draft and edit the text. The authors designed and ran the experiment, reviewed the analyses "
              "and text, and take responsibility for the content.]")

    # -------------------------------------------------------------- REFERENCES
    doc.add_paragraph("REFERENCES", style=HREF)
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph(style=REF)
        p.add_run(f"[{i}] {ref}")


def add_table_borders(table, size="4"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def add_algorithm_box(doc):
    lines = [
        "Require: block values y_{l,k} (levels l = 1..L, blocks k = 1..5); threshold φ; draws B",
        "1:  ȳ_l ← (1/5) Σ_k y_{l,k};  p ← argmax_l ȳ_l",
        "2:  j ← min{l > p : ȳ_l ≤ φ ȳ_p}          ▷ bracket [Δθ_{j-1}, Δθ_j]",
        "3:  x̂_M4 ← root of PCHIP(Δθ_l, ȳ_l) = φ ȳ_p in the bracket",
        "4:  x̂_M3 ← linear root in the bracket; x̂_m ← crossing of the fitted curve, m ∈ {M1, M2, M5}",
        "5:  for r = 1 to B do",
        "6:      resample the five blocks of each level with replacement; repeat steps 1–4 → x_m^(r)",
        "7:  end for",
        "8:  return bracket; x̂_M4 with the 2.5 and 97.5 percentiles of x_M4^(r); [min_m x̂_m, max_m x̂_m]",
    ]
    box = doc.add_table(rows=1, cols=1)
    add_table_borders(box)
    cell = box.rows[0].cells[0]
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        run = par.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    doc.add_paragraph()


def add_table_divider(doc):
    headers = ["Zone", "d (m)", "Δθ = 4°", "Δθ = 10°"]
    rows = [
        ["16×16", "0.3", "−1.3 (−1.8, −0.8)", "−0.7 (−1.3, −0.1)"],
        ["16×16", "0.6", "+4.1 (+3.5, +4.8)", "+3.9 (+2.8, +5.0)"],
        ["16×16", "0.9", "+1.6 (+1.1, +2.2)", "+2.5 (+1.1, +4.0)"],
        ["4×4", "0.3", "−3.9 (−4.5, −3.3)", "−7.7 (−11.8, −4.0)"],
        ["4×4", "0.6", "+4.2 (+2.7, +5.7)", "+17.2 (+13.0, +22.0)"],
        ["4×4", "0.9", "+0.2 (−1.1, +1.5)", "−0.1 (−2.8, +2.7)"],
    ]
    t = doc.add_table(rows=1, cols=len(headers))
    add_table_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8)
                r.font.name = "Times New Roman"
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.name = "Times New Roman"
    doc.add_paragraph()


def add_table_crossover(doc):
    headers = ["Zone", "d (m)", "Bracket", "Crossover\n(PCHIP)", "Range of\nthree est.",
               "1\u2009% angle", "Model,\nas coded", "Model,\nF in-plane"]
    rows = [
        ["16\u00d716", "0.3", "16\u201318", "16.15", "16.15\u201316.20", "19.9", "23.4", "19.5"],
        ["16\u00d716", "0.6", "16\u201318", "16.97", "16.97\u201317.13", "20.9", "25.2", "21.3"],
        ["16\u00d716", "0.9", "16\u201318", "17.90", "17.90\u201317.94", "21.6", "25.8", "21.9"],
        ["4\u00d74", "0.3", "8\u201310", "8.22", "8.20\u20138.33", "11.4", "17.3", "15.2"],
        ["4\u00d74", "0.6", "8\u201310", "9.47", "9.36\u20139.66", "12.5", "19.2", "17.1"],
        ["4\u00d74", "0.9", "10\u201312", "10.80", "10.80\u201311.02", "13.1", "19.8", "17.7"],
    ]
    t = doc.add_table(rows=1, cols=len(headers))
    add_table_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8)
                r.font.name = "Times New Roman"
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.name = "Times New Roman"
    for row in t.rows:
        for cell, w in zip(row.cells, (0.7, 0.5, 0.7, 0.9, 1.5, 0.8, 1.0, 1.0)):
            cell.width = Inches(w * LAYOUT["table_w"] / 7.1)
    note = doc.add_paragraph(style="Caption L-SENS")
    note.add_run("Bracket and range (linear, PCHIP and tilted logistic) as defined in Methodology; bootstrap "
                 "95\u2009% half-widths of the PCHIP value are 0.01\u20130.07\u00b0 (within-session only).")


def build():
    # Normally the downloaded template is the base. The template file was deleted from the project folder on
    # 11 Oct 2026 (it is in the Recycle Bin), so when it is absent the previously filled output is the base:
    # it carries every template style, header and section setting, and only the body (from INTRODUCTION on)
    # is rebuilt; the front matter is already filled and is left alone.
    if TEMPLATE.exists():
        doc = Document(str(TEMPLATE))
        fresh = True
    elif OUT.exists():
        doc = Document(str(OUT))
        fresh = False
        print(f"template not found; rebuilding the body of {OUT.name} in place")
    else:
        raise SystemExit(f"neither {TEMPLATE.name} nor {OUT.name} found")
    styles = {s.name: s for s in doc.styles}

    if fresh:
        fill_front_matter(doc)
    # the template's body section is flagged "different first page"; the body now starts mid-page 1 and is split
    # into several continuous sections below, so only the front-matter section keeps that flag
    sp = doc.element.body.find(qn("w:sectPr"))
    tp = sp.find(qn("w:titlePg"))
    if tp is not None:
        sp.remove(tp)
    clear_body(doc)
    build_body(doc, styles)

    cp = doc.core_properties
    cp.author = "Jessica"
    cp.last_modified_by = "Jessica"
    cp.title = "Angle-Dependent Crosstalk Between Two Adjacent VL53L1X Time-of-Flight Sensors"
    cp.subject = ""
    cp.comments = ""
    import datetime
    now = datetime.datetime(2026, 10, 11, 15, 0, 0, tzinfo=datetime.timezone.utc)
    cp.created = now
    cp.modified = now
    cp.revision = 1

    doc.save(str(OUT))
    scrub_app_xml(OUT)
    prune_orphaned_media(OUT)
    return OUT


def prune_orphaned_media(path: Path):
    """clear_body() deletes the template's own sample paragraphs, including its one embedded Equation-Editor
    object (an OLE object with .emf/.wmf preview images); python-docx does not garbage-collect the now-unused
    media/relationship entries that referenced it. Remove them so the file carries nothing but what it uses."""
    import zipfile
    import tempfile
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    doc_xml = data["word/document.xml"].decode("utf-8")
    rels_path = "word/_rels/document.xml.rels"
    rels_xml = data[rels_path].decode("utf-8")
    kept_rels = []
    dropped_targets = []
    for m in re.finditer(r'<Relationship [^>]*/>', rels_xml):
        tag = m.group(0)
        rid = re.search(r'Id="(rId\d+)"', tag).group(1)
        tgt = re.search(r'Target="([^"]+)"', tag)
        is_media_or_object = tgt and ("media/" in tgt.group(1) or "embeddings/" in tgt.group(1))
        if is_media_or_object and f'"{rid}"' not in doc_xml:
            dropped_targets.append("word/" + tgt.group(1))
            continue
        kept_rels.append(tag)
    if dropped_targets:
        header = re.match(r'^.*?<Relationships[^>]*>', rels_xml, re.S).group(0)
        rels_xml = header + "".join(kept_rels) + "</Relationships>"
        data[rels_path] = rels_xml.encode("utf-8")
        for t in dropped_targets:
            data.pop(t, None)
    tmp = Path(tempfile.mktemp(suffix=".docx"))
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n in data:
                zout.writestr(n, data[n])
    shutil.move(str(tmp), str(path))
    return dropped_targets


def scrub_app_xml(path: Path):
    import zipfile
    import tempfile
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    if "docProps/app.xml" in data:
        app = data["docProps/app.xml"].decode("utf-8", errors="ignore")
        app = re.sub(r"<Application>.*?</Application>", "<Application></Application>", app)
        data["docProps/app.xml"] = app.encode("utf-8")
    tmp = Path(tempfile.mktemp(suffix=".docx"))
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(str(tmp), str(path))


if __name__ == "__main__":
    out = build()
    print("wrote", out)
