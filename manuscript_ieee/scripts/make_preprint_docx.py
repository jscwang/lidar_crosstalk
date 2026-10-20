"""Word copy of the comprehensive (preprint) version, for mark-up.

Plain single-column document, Times New Roman, no journal header (the journal template's masthead would present a
preprint as an IEEE Sensors Letters page). The text is the one in fill_sensors_letters_docx.py, switched to its
preprint layout (LAYOUT["preprint"]): the same body as the letter plus the supplement's content merged in
(estimator labels M1-M5, Algorithm 1, Fig. 3, Table II), and "This paper" for "This letter". It matches
arxiv/paper.tex, so comments made on this file map one-to-one onto that PDF.

Run:   python manuscript_ieee/scripts/make_preprint_docx.py
Out:   manuscript_ieee/arxiv/paper.docx

Author: Jessica
"""
from __future__ import annotations

import datetime
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import fill_sensors_letters_docx as F  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "arxiv" / "paper.docx"
TITLE = "Angle-Dependent Crosstalk Between Two Adjacent VL53L1X Time-of-Flight Sensors"

ABSTRACT = (
    "Time-of-flight ranging sensors mounted side by side can disturb one another, but we found no published "
    "measurement of how this depends on the angle between them. We measured the extra background light that "
    "appears in one VL53L1X sensor when a second, identical sensor 20 millimeters away is also ranging, as the "
    "two were turned apart by up to 24 degrees. One pair of sensors was tested against a grey board at 0.3, 0.6 "
    "and 0.9 meters in two window sizes, with five repeats per setting and a divider test. The extra light fell "
    "to 10 percent of its peak between 16 and 18 degrees for the full array and between 8 and 12 degrees for the "
    "window, and to 1 percent at about 20 to 22 and 11 to 13 degrees; these angles grew with distance. Methods "
    "for locating the 10 percent angle disagreed by up to one degree, ten to forty-four times the sampling "
    "uncertainty, so we report a bracket and an interpolated value. A geometric model using the datasheet field "
    "of view overestimated the angle at which the effect disappears by three to seven degrees, and no single "
    "emitter beam angle fitted both windows. A thin divider changed the effect by no more than about four "
    "percent in ten of twelve settings, which does not show that direct coupling is absent. Range readings were "
    "barely affected. Turning each sensor outward by about 11 degrees reduced the effect below 1 percent of its "
    "peak for the full array."
)


def add_style(doc, name, size=11, bold=False, italic=False, align=None, after=6, left=None, first=None,
              base="Normal"):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = italic
    st.font.color.rgb = RGBColor(0, 0, 0)
    pf = st.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Inches(left)
    if first is not None:
        pf.first_line_indent = Inches(first)
    return st


def setup(doc):
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(11)
    for name, size, italic in (("Heading 1", 13, False), ("Heading 2", 11.5, True)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = not italic
        st.font.italic = italic
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12 if not italic else 8)
        st.paragraph_format.space_after = Pt(4)
    # the paragraph styles that fill_sensors_letters_docx.build_body asks for by their template names
    add_style(doc, "Text L-SENS", 11, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_style(doc, "Caption L-SENS", 9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
    add_style(doc, "Acknowledgment L-SENS", 11, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_style(doc, "References L-SENS", 10, left=0.35, first=-0.35, after=4)
    h = add_style(doc, "Heading Refs L-SENS", 13, bold=True, after=4)
    h.paragraph_format.space_before = Pt(12)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)


def build():
    # single-column geometry (text width 6.5 in) instead of the journal's two columns
    F.LAYOUT.update({"two_col": False, "preprint": True, "col_w": 6.5, "full_w": 6.4, "fig1_w": 3.6,
                     "table_w": 6.4})
    doc = Document()
    setup(doc)
    styles = {s.name: s for s in doc.styles}

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(17)
    a = doc.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.add_run("[Author names and affiliations to be supplied]")
    n = doc.add_paragraph()
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = n.add_run("Preprint, not peer reviewed. Word copy for comments; the typeset version is paper.pdf.")
    r.italic = True
    r.font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract—")
    r.bold = True
    r.italic = True
    p.add_run(ABSTRACT)
    p = doc.add_paragraph()
    r = p.add_run("Index Terms—")
    r.bold = True
    r.italic = True
    p.add_run("Crosstalk, interference, sensor arrays, time-of-flight sensors, VL53L1X.")

    F.build_body(doc, styles)

    cp = doc.core_properties
    cp.author = "Jessica"
    cp.last_modified_by = "Jessica"
    cp.title = TITLE + " (preprint)"
    cp.comments = ""
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
    cp.created = now
    cp.modified = now
    cp.revision = 1

    doc.save(str(OUT))
    F.scrub_app_xml(OUT)
    return OUT


if __name__ == "__main__":
    print("wrote", build())
