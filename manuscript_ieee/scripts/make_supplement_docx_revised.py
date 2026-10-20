"""Word version of the REVISED supplementary material (supplementary/supplement.tex, after Jessica's comments,
Comments v1.docx, 14-15 Oct 2026). Adds the new "Acceptance rules" section (shown in red, like the rest of
this revision) and updates the intro paragraph; everything else is unchanged from make_supplement_docx.py.

Run:   python manuscript_ieee/scripts/make_supplement_docx_revised.py
Out:   manuscript_ieee/supplementary/supplement_revised.docx

Author: Jessica
"""
from __future__ import annotations

import datetime
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import fill_sensors_letters_docx as F  # noqa: E402
import make_revised_docx as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
OUT = ROOT / "supplementary" / "supplement_revised.docx"
TITLE = "Angle-Dependent Crosstalk Between Two Adjacent VL53L1X Time-of-Flight Sensors"
RED = RGBColor(0xC0, 0x00, 0x00)


def setup(doc):
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(11)
    n.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 13), ("Title", 16)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("List Bullet",):
        doc.styles[name].font.name = "Times New Roman"
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)


def para(doc, latex, style=None, align=None, size=None, red=False):
    p = doc.add_paragraph(style=style)
    R.add_rich_tracked(p, latex, force_red=red)
    if align is not None:
        p.alignment = align
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    return p


def caption(doc, latex, red=False):
    return para(doc, latex, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=9.5, red=red)


def algorithm_box(doc):
    p = doc.add_paragraph()
    r = p.add_run("Algorithm S1")
    r.bold = True
    p.add_run("  Crossover for one cell (zone, distance)")
    lines = [
        "Require: block values y_{l,k} (levels l = 1..L, blocks k = 1..5); threshold φ; draws B",
        "1:  ȳ_l ← (1/5) Σ_k y_{l,k};  p ← argmax_l ȳ_l",
        "2:  j ← min{l > p : ȳ_l ≤ φ ȳ_p}          ▷ bracket [Δθ_{j-1}, Δθ_j]",
        "3:  x̂_M4 ← root of PCHIP(Δθ_l, ȳ_l) = φ ȳ_p in the bracket",
        "4:  x̂_M3 ← linear root in the bracket; x̂_m ← crossing of the fitted curve, m ∈ {M1, M2, M5}",
        "5:  for r = 1 to B do",
        "6:      resample the five blocks of each level with replacement; repeat steps 1–4 → x_m^(r)",
        "7:  end for",
        "8:  return bracket; x̂_M4 with the 2.5 and 97.5 percentiles of x_M4^(r); [min_m x̂_m, max_m x̂_m]",
    ]
    box = doc.add_table(rows=1, cols=1)
    F.add_table_borders(box)
    cell = box.rows[0].cells[0]
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        run = par.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    doc.add_paragraph()


def table_divider(doc):
    headers = ["Zone", "d (m)", "Δθ = 4°", "Δθ = 10°"]
    rows = [
        ["16×16", "0.3", "−1.3 (−1.8, −0.8)", "−0.7 (−1.3, −0.1)"],
        ["16×16", "0.6", "+4.1 (+3.5, +4.8)", "+3.9 (+2.8, +5.0)"],
        ["16×16", "0.9", "+1.6 (+1.1, +2.2)", "+2.5 (+1.1, +4.0)"],
        ["4×4", "0.3", "−3.9 (−4.5, −3.3)", "−7.7 (−11.8, −4.0)"],
        ["4×4", "0.6", "+4.2 (+2.7, +5.7)", "+17.2 (+13.0, +22.0)"],
        ["4×4", "0.9", "+0.2 (−1.1, +1.5)", "−0.1 (−2.8, +2.7)"],
    ]
    t = doc.add_table(rows=1, cols=len(headers))
    F.add_table_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def build():
    doc = Document()
    setup(doc)

    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run("Supplementary Material")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(TITLE)
    r.font.size = Pt(13)
    d = doc.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_run("Draft, revised following Jessica's comments — new text shown in red").italic = True

    para(doc, r"This document accompanies the letter of the same title. It holds material moved out of the "
              r"letter to keep it short: the procedure used to locate the crossover angle (Algorithm~S1), "
              r"\new{the full block- and run-acceptance tolerances,} the divider results (Table~S1) and a "
              r"figure comparing the crossover estimators in all six cells (Fig.~S1). Every number here is "
              r"computed from the data files and scripts in the accompanying repository.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_heading("1  Crossover estimators and procedure", level=1)
    para(doc, "For each cell (zone and distance) the extra ambient rate is the mean, over sensors A and B, of "
              "the ambient rate in the A+B run minus the ambient rate in that sensor's solo run, one value per "
              "block. Each cell has five blocks at each $\\dth$. The crossover is the angle at which the "
              "response falls to a fraction $\\varphi$ of its reference level; the letter reports "
              "$\\varphi=0.10$ and $\\varphi=0.01$. Five estimators were compared:",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for b in [
        "M1: three-parameter logistic fitted to all levels; reference is the fitted value at $\\dth=0$;",
        "M2: the same logistic fitted only from the observed peak level outward; reference is the fitted value "
        "there (identical to M1 when the peak is at $\\dth=0$);",
        "M3: piecewise-linear interpolation of the block means; reference is the observed peak mean; first "
        "crossing after the peak;",
        "M4: shape-preserving cubic (PCHIP) interpolation of the block means, same reference and crossing rule "
        "(the primary estimator in the letter);",
        "M5: four-parameter tilted logistic $a(1+b\\,\\dth)/(1+\\exp[(\\dth-x_0)/w])$, which may rise before it "
        "falls; reference is its own maximum; first crossing after that maximum.",
    ]:
        para(doc, b, style="List Bullet")
    para(doc, "Uncertainty is a stratified block bootstrap (the five blocks of each level resampled with "
              "replacement, 2000 draws, one common set of draws for all estimators, percentile intervals at "
              "95\\,\\%). Algorithm~S1 gives the procedure for one cell and one threshold.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    algorithm_box(doc)
    para(doc, "In steps 3–4 the references of M1, M2 and M5 follow their definitions above rather than "
              "the observed peak $\\bar y_p$.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_heading("2  Acceptance rules", level=1)
    para(doc, r"\new{A run was accepted if it had no abort or I$^2$C/read error, the idle sensor stayed "
              r"silent, the reading count was within \SI{5}{\percent} of expected and at most "
              r"\SI{1}{\percent} of results were missed. A block was accepted if its runs were and each "
              r"sensor's solo mean ambient rate was within \SI{10}{\percent} (at least \SI{0.5}{kcps}) of "
              r"that of the previous accepted block at the same distance, angle and zone; failures were "
              r"repeated up to three times. Blocks failing three times (1, 6 and 7 of 120 at 0.3, 0.6 and "
              r"\SI{0.9}{\meter}) were re-measured in separate refill sessions and replaced.}",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_heading("3  Divider results", level=1)
    para(doc, "Table~S1 gives the change in the extra ambient rate when the thin divider stood between the "
              "sensor windows, at the two angles where it was re-measured. The divider was confounded with "
              "session, and the strip's size and position were not recorded, so the table cannot show that "
              "direct coupling is absent.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    caption(doc, "Table S1. Change in the extra ambient rate with the divider, \\si{\\percent} of the "
                 "no-divider cell mean (\\SI{95}{\\percent} bootstrap interval)")
    table_divider(doc)

    doc.add_heading("4  Comparison of the estimators", level=1)
    para(doc, "Figure~S1 shows the block means, three of the five estimators and the 10\\,\\% crossing each "
              "one gives, in every cell. In the $4\\times4$ zone at 0.6 and 0.9\\,m the response peaks at "
              "$\\dth=4^\\circ$, not at $0^\\circ$; the three curves then differ most at small angles, near "
              "the peak, while their 10\\,\\% crossings stay close.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(FIG / "fig_crossover_methods.png"), width=Inches(6.5))
    caption(doc, "Fig. S1. Extra ambient rate against $\\dth$ in each cell, as a percentage of the largest "
                 "cell mean. Points are block means ± SD ($n=5$). Curves: M1 logistic, M5 tilted logistic "
                 "and M4 PCHIP interpolant. Triangles mark where each of these three crosses 10\\,\\% of the "
                 "peak (colour identifies the estimator). The dash-dotted vertical line is the geometric "
                 "prediction of the letter's overlap model with the datasheet field of view; the sensitivity "
                 "variants are in the repository.")

    cp = doc.core_properties
    cp.author = "Jessica"
    cp.last_modified_by = "Jessica"
    cp.title = "Supplementary material (revised): " + TITLE
    cp.comments = ""
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
    cp.created = now
    cp.modified = now
    cp.revision = 1

    doc.save(str(OUT))
    F.scrub_app_xml(OUT)
    return OUT


if __name__ == "__main__":
    print("wrote", build())
